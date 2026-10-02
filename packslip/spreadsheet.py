"""Read the weekly .xlsx order file into headers + one dict per customer row."""

from __future__ import annotations

import csv
import datetime as dt
import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException


_TOTAL_LABELS = {"total", "totals", "grand total", "subtotal", "sum"}


class SpreadsheetError(Exception):
    """A problem with the spreadsheet, worded for a non-technical reader."""


@dataclass
class Sheet:
    path: Path
    sheet_name: str
    sheet_names: list[str]
    headers: list[str]
    rows: list[dict[str, str]] = field(default_factory=list)
    header_row: int = 1  # 1-based row number in Excel

    @property
    def customer_count(self) -> int:
        return len(self.rows)


def format_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return f"{value:%b} {value.day}, {value.year}"
        hour = value.hour % 12 or 12
        return f"{value:%b} {value.day}, {value.year} {hour}:{value:%M} {value:%p}"
    if isinstance(value, dt.date):
        return f"{value:%b} {value.day}, {value.year}"
    if isinstance(value, dt.time):
        hour = value.hour % 12 or 12
        return f"{hour}:{value:%M} {value:%p}"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _open(path: Path):
    if not path.exists():
        raise SpreadsheetError(
            f"The file “{path.name}” couldn't be found. It may have been moved or renamed."
        )
    suffix = path.suffix.lower()
    if path.name.startswith("~$"):
        raise SpreadsheetError(
            "That's a temporary file Excel creates while a spreadsheet is open. "
            "Please choose the real spreadsheet instead."
        )
    if suffix == ".xls":
        raise SpreadsheetError(
            "This spreadsheet is in the old Excel format (.xls).\n\n"
            "Open it in Excel or Numbers and use File → Save As (or Export) to save it "
            "as an Excel Workbook (.xlsx), then try again."
        )
    if suffix in (".numbers", ".ods"):
        raise SpreadsheetError(
            f"This app reads Excel (.xlsx) and .csv files, but this file is a {suffix} file.\n\n"
            "Open it in Numbers or Excel and export it as an Excel Workbook (.xlsx), then try again."
        )
    if suffix not in (".xlsx", ".xlsm", ".csv"):
        raise SpreadsheetError(
            f"“{path.name}” isn't a spreadsheet this app can read. Please choose a file ending in .xlsx or .csv."
        )
    if suffix == ".csv":
        return None
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise SpreadsheetError(
            f"“{path.name}” couldn't be opened. It may be damaged, password-protected, "
            "or still being saved.\n\nTry opening it in Excel, saving it again, and re-loading it."
        ) from None


def list_sheets(path: str | Path) -> list[str]:
    wb = _open(Path(path))
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _find_header_row(raw: list[tuple]) -> int:
    """Index of the header row: the first row that is 'about as wide' as the table.

    Handles a title line (e.g. "Orders – week of Sep 30") above the real headers.
    """
    counts = [sum(1 for v in r if format_value(v)) for r in raw[:20]]
    widest = max(counts, default=0)
    if widest == 0:
        return -1
    for i, c in enumerate(counts):
        if c >= max(2, widest * 0.6) or (widest == 1 and c == 1):
            return i
    return counts.index(widest)


def _read_csv(path: Path) -> list[tuple]:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if "\x00" in text:
        raise SpreadsheetError(f"“{path.name}” doesn't look like a text .csv file. Try exporting it again.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [tuple(r) for r in csv.reader(io.StringIO(text), dialect)]


def load(path: str | Path, sheet_name: str | None = None) -> Sheet:
    path = Path(path)
    wb = _open(path)
    if wb is None:  # .csv
        raw, title, names = _read_csv(path), path.stem, [path.stem]
    else:
        raw, title, names = _read_workbook(wb, sheet_name)
    return _build(path, raw, title, names)


def _read_workbook(wb, sheet_name):
    try:
        names = list(wb.sheetnames)
        if sheet_name and sheet_name in names:
            ws = wb[sheet_name]
        else:
            ws = wb.worksheets[0]
            # Prefer the first sheet that actually has data.
            for candidate in wb.worksheets:
                if any(any(v is not None for v in r) for r in candidate.iter_rows(max_row=20, values_only=True)):
                    ws = candidate
                    break
        raw = [tuple(r) for r in ws.iter_rows(values_only=True)]
        title = ws.title
    finally:
        wb.close()
    return raw, title, names


def _build(path: Path, raw: list[tuple], title: str, names: list[str]) -> Sheet:
    hidx = _find_header_row(raw)
    if hidx < 0:
        raise SpreadsheetError(
            f"The sheet “{title}” in “{path.name}” is empty. "
            "Make sure you picked this week's order spreadsheet."
        )

    header_cells = [format_value(v) for v in raw[hidx]]
    # Trim trailing blank columns (Excel often reports extra empty ones).
    width = max((i + 1 for i, h in enumerate(header_cells) if h), default=0)
    for r in raw[hidx + 1:]:
        for i in range(len(r) - 1, width - 1, -1):
            if format_value(r[i]):
                width = max(width, i + 1)
                break

    headers: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        h = header_cells[i] if i < len(header_cells) else ""
        h = " ".join(h.split()) or f"Column {get_column_letter(i + 1)}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        headers.append(h)

    rows = []
    for r in raw[hidx + 1:]:
        values = [format_value(r[i]) if i < len(r) else "" for i in range(width)]
        if not any(values):
            continue
        if any(v.strip().lower().rstrip(":") in _TOTAL_LABELS for v in values):
            continue  # a summary line like "TOTALS", not a customer
        rows.append(dict(zip(headers, values)))

    if not rows:
        raise SpreadsheetError(
            f"“{path.name}” has column headings but no customer rows underneath them."
        )
    return Sheet(path=path, sheet_name=title, sheet_names=names, headers=headers,
                 rows=rows, header_row=hidx + 1)
