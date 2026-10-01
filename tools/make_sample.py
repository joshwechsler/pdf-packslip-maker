"""Create sample/sample_orders.xlsx: fake data until the Owner supplies a real file (SRS §8).

Usage:  python tools/make_sample.py [rows] [out.xlsx]
"""

import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

FIRST = ["Jane", "Marcus", "Priya", "Tom", "Aisha", "Luis", "Chen", "Hannah", "Omar", "Grace", "Diego", "Maya"]
LAST = ["Smith", "Johnson", "Patel", "Nguyen", "Garcia", "Okafor", "Kim", "Rossi", "Cohen", "Brown"]
STREETS = ["Maple Ave", "Oak St", "Elm Dr", "Cedar Ln", "Lakeview Rd", "Hillcrest Blvd"]
ITEMS = ["Chicken Bowl", "Greek Salad", "Beef Chili", "Veggie Wrap", "Salmon & Rice",
         "Turkey Meatballs", "Banana Bread", "Overnight Oats", "Lentil Soup", "Protein Box"]
NOTES = ["", "", "Leave at side door", "No nuts please", "Call on arrival", "Gate code 4412", ""]


def make(rows: int, out: Path, seed: int = 7) -> Path:
    rnd = random.Random(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Weekly orders – week of Sep 28, 2026"])
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["Order No", "Customer", "Delivery Date", "Delivery Address", "Phone",
               "Order Items", "Qty", "Special Instructions"]
    ws.append(headers)
    for c in ws[2]:
        c.font = Font(bold=True)
    for i in range(rows):
        picks = rnd.sample(ITEMS, rnd.randint(1, 5))
        ws.append([
            1001 + i,
            f"{rnd.choice(FIRST)} {rnd.choice(LAST)}",
            dt.datetime(2026, 10, 1) + dt.timedelta(days=rnd.randint(0, 2)),
            f"{rnd.randint(10, 999)} {rnd.choice(STREETS)}\nSpringfield, IL 627{rnd.randint(0, 99):02d}",
            f"(555) {rnd.randint(200, 999)}-{rnd.randint(1000, 9999)}",
            "\n".join(picks),
            "\n".join(str(rnd.randint(1, 4)) for _ in picks),
            rnd.choice(NOTES),
        ])
    for col, width in zip("ABCDEFGH", (10, 20, 14, 30, 16, 26, 6, 26)):
        ws.column_dimensions[col].width = width
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 12
    path = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / "sample" / "sample_orders.xlsx"
    print(make(n, path))
